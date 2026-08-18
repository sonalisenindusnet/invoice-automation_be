"""
entity_resolver.py

Decides WHICH tab (USA / UK / Poland 2026 — the only tabs this is wired up
to use right now) a new invoice request belongs in, based on the client
company name in the CP's "Invoice Summary" email. A CP email can be for a
client billed by any of your entities, so we can't just assume everything
goes into one tab.

Two-step lookup, in order:
  1. config/company_tab_map.json — the file you maintain by hand. Exact
     match (case/spacing/punctuation-insensitive) against "company_name".
  2. Fallback: scan every known tab's company-name column in the actual
     tracker for an existing row with that exact company name. If found,
     the match is written BACK into company_tab_map.json automatically
     (with a note that it was auto-detected), so the next email for that
     same client resolves instantly from step 1.

If neither step finds a match, this returns None — the caller must NOT
guess a tab for money. See append_invoice_to_excel.py's "entity_unresolved"
status.

Note: India and Kar Ventures are intentionally NOT in ENTITY_SCHEMA_PATHS
below — those are GST/India-specific and out of scope for now (foreign
clients only). A client that resolves to one of those would come back
unresolved rather than routed anywhere. Tell me if/when you want India
support added back — it's a small change (add the tab back to this dict).

---

**SUPERSEDED for live routing, 2026-08-13.** Per explicit instruction, tab
selection is no longer decided by company name at all: "the account team
will put from which tab (eg. india/uk/usa/poland etc) the invoice will be
generated. this tab selection will not depending on the MIS API response
or any other parameter which I uses earlier." `resolve_entity_from_email_text()`
below is what the live server (`email_server.py`) actually calls now.
`resolve_entity_for_company()` and `company_tab_map.json` above are left in
place, untouched, but are no longer called from the live intake path —
kept only in case this ever needs to come back.

Format (tightened same day, after "let then free make trouble for us" —
an earlier version scanned the whole body for any "<name> tab" sentence,
which risked misfiring on quoted/forwarded history): the account team must
put, as the very FIRST line of the email, exactly:
    Tab: UK
(or "Tab: USA" / "Tab: Poland" — case-insensitive, flexible spacing around
the colon). Nothing else anywhere else in the email is scanned for this.
"""
import json
import re
from datetime import date
from pathlib import Path

import openpyxl

# 2026-08-14: see tracker_io.py's module docstring -- transparently opens
# either a local .xlsx path (unchanged) or the live Google Sheet, depending
# on what's passed in. resolve_entity_for_company() below is superseded by
# resolve_entity_from_email_text() for live routing (see this module's
# docstring), so this swap is for consistency/future-proofing only.
from utils.tracker_io import load_tracker_with_retry as load_workbook_with_retry

CONFIG_DIR = Path(__file__).resolve().parent.parent.parent / "config"
COMPANY_MAP_PATH = CONFIG_DIR / "company_tab_map.json"
TABS_DIR = CONFIG_DIR / "tabs"

ENTITY_SCHEMA_PATHS = {
    "usa": TABS_DIR / "usa.json",
    "uk": TABS_DIR / "uk.json",
    "poland": TABS_DIR / "poland.json",
    # "india" / "kar_ventures" / "singapore" deliberately absent — foreign
    # clients only, for now.
}

_LEGAL_SUFFIXES = [
    "private limited", "pvt ltd", "pte ltd", "sp z o o", "sarl", "llc",
    "limited", "ltd", "inc", "corp", "corporation", "co",
]


def load_schema(entity_key):
    path = ENTITY_SCHEMA_PATHS.get(entity_key)
    if path is None or not path.exists():
        raise ValueError(f"No Excel tab schema configured for entity '{entity_key}'")
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _load_company_map():
    if not COMPANY_MAP_PATH.exists():
        return {"companies": [], "entity_tabs": {}}
    with open(COMPANY_MAP_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def _save_company_map(company_map):
    with open(COMPANY_MAP_PATH, "w", encoding="utf-8") as f:
        json.dump(company_map, f, indent=2, ensure_ascii=False)
        f.write("\n")


def _normalize(name):
    """Lowercase, strip punctuation, collapse whitespace, and drop a trailing
    legal-entity suffix — so 'Accuride International Ltd.' and
    'accuride international ltd' resolve to the same key, without being so
    loose it risks matching the wrong client."""
    if not name:
        return ""
    n = name.lower().strip()
    n = re.sub(r"[.,()]", "", n)
    n = re.sub(r"\s+", " ", n).strip()
    for suffix in _LEGAL_SUFFIXES:
        pattern = r"\s+" + re.escape(suffix) + r"$"
        n = re.sub(pattern, "", n)
    return n.strip()


def resolve_entity_for_company(company_name, xlsx_path=None):
    """
    Returns (entity_key, resolution) where resolution is one of:
      "mapped"      — found in company_tab_map.json
      "auto_detected" — found by scanning an existing tab; also just got
                        added to company_tab_map.json for next time
      None entity_key, resolution "unresolved" — not found anywhere;
                        caller must flag this, never guess.
    """
    if not company_name or not company_name.strip():
        return None, "unresolved"

    target = _normalize(company_name)
    company_map = _load_company_map()

    for entry in company_map.get("companies", []):
        if _normalize(entry.get("company_name", "")) == target:
            return entry["entity_key"], "mapped"

    if xlsx_path is None:
        return None, "unresolved"

    try:
        wb = load_workbook_with_retry(xlsx_path, data_only=True)
    except FileNotFoundError:
        return None, "unresolved"

    for entity_key, schema_path in ENTITY_SCHEMA_PATHS.items():
        if not schema_path.exists():
            continue
        with open(schema_path, "r", encoding="utf-8") as f:
            schema = json.load(f)
        sheet_name = schema["sheet_name"]
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        cols = schema["columns"]
        company_key = schema.get("company_column_key", "client_company")
        col_idx = next((i for i, c in enumerate(cols) if c["key"] == company_key), None)
        if col_idx is None:
            continue
        start = schema.get("header_row_index", 1) + 1
        for row in ws.iter_rows(min_row=start, max_row=ws.max_row, values_only=True):
            if row is None or col_idx >= len(row):
                continue
            existing_name = row[col_idx]
            if existing_name and _normalize(str(existing_name)) == target:
                company_map.setdefault("companies", []).append({
                    "company_name": company_name.strip(),
                    "entity_key": entity_key,
                    "note": f"auto-detected from an existing '{sheet_name}' row on {date.today().isoformat()}",
                })
                _save_company_map(company_map)
                return entity_key, "auto_detected"

    return None, "unresolved"


# STRICT FORMAT, tightened 2026-08-13. First version of this scanned the
# whole raw body for any "<name> tab" sentence anywhere -- flagged as too
# loose ("isnot there is a special format, let then free make trouble for
# us"): quoted/forwarded history further down an email, or a stray
# coincidental mention, could misfire. Tightened per instruction ("is it
# better to put it in the first line like Tab:<tab>") to a single fixed
# line, in the exact form "Tab: <name>", and it MUST be the first non-blank
# line of the email -- nothing else anywhere in the body is scanned or
# accepted anymore. This also eliminates the old "ambiguous" case entirely
# (there's only ever one line being checked now, so there's nothing left to
# be ambiguous between).
_TAB_LABEL_RE = re.compile(r"^\s*tab\s*:\s*(usa|uk|poland|india)\s*$", re.IGNORECASE)


def resolve_entity_from_email_text(raw_text):
    """
    NEW 2026-08-13, per explicit instruction: which tab an invoice goes
    into is now decided ENTIRELY by an explicit instruction the account
    team writes into the email -- NOT by the client company name, NOT by
    company_tab_map.json, NOT by the MIS API response, and not by any
    other parameter used before this. This is what the live server calls
    now instead of resolve_entity_for_company() -- see the module
    docstring above.

    Required format: the FIRST non-blank line of the email must be exactly
        Tab: UK
    (or "Tab: USA" / "Tab: Poland" -- case-insensitive, flexible spacing
    around the colon, e.g. "Tab:UK" or "Tab : UK" both also work). Nothing
    else in the email is scanned for this -- a tab name mentioned anywhere
    other than that first line has no effect.

    Returns (entity_key, resolution, message):
      (entity_key, "explicit_from_email", None)
          -- the first line was exactly "Tab: usa/uk/poland". Use it
             directly, no further lookup.
      (None, "india_out_of_scope", message)
          -- first line was "Tab: India", but per explicit instruction
             (2026-08-13: "India tab will be developed later, for now
             keep it out of scope") India is not wired into live routing
             yet. Never silently routed anywhere -- caller must flag this
             for manual handling.
      (None, "no_tab_instruction_found", message)
          -- the first non-blank line didn't match "Tab: <name>" at all
             (missing, malformed, or in the wrong place). Company-name
             auto-detection is intentionally NOT used as a fallback here
             anymore -- caller must flag this too.
    """
    if not raw_text or not raw_text.strip():
        return None, "no_tab_instruction_found", (
            "Email body is empty -- no tab instruction found (expected the first "
            "line to read exactly 'Tab: UK', 'Tab: USA', or 'Tab: Poland')."
        )

    first_line = next((line for line in raw_text.splitlines() if line.strip()), "")
    match = _TAB_LABEL_RE.match(first_line)

    if not match:
        return None, "no_tab_instruction_found", (
            f"The first line of the email must be exactly 'Tab: <name>' (e.g. "
            f"'Tab: UK') -- got {first_line.strip()!r} instead. Company-name "
            f"auto-detection is no longer used for routing, so this can't be "
            f"resolved automatically."
        )

    entity_key = match.group(1).lower()
    if entity_key == "india":
        return None, "india_out_of_scope", (
            "First line says 'Tab: India', but India is not wired into live "
            "routing yet (GST rule and our own company details are still "
            "pending) -- needs manual handling."
        )

    return entity_key, "explicit_from_email", None
